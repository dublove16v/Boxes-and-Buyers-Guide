import csv
import re
from datetime import datetime
from io import BytesIO, StringIO

from openpyxl import load_workbook

DAY_BANNER = re.compile(r"^(MON|TUE|TUES|WED|WEDS|THU|THUR|THURS|FRI|SAT|SUN)(DAY|S)?$", re.I)
WEEK_NAME = re.compile(r"BOXES\s+(\d{1,2})-(\d{1,2})-(\d{2,4})", re.I)
SHEET_DATE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")
DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def week_date_from_name(name: str) -> str:
    match = WEEK_NAME.search(name)
    if not match:
        return ""
    month, day, year = (int(match.group(1)), int(match.group(2)), int(match.group(3)))
    if year < 100:
        year += 2000
    try:
        return datetime(year, month, day).date().isoformat()
    except ValueError:
        return ""


def _header_index(headers: list[str], name: str) -> int:
    target = name.upper()
    for index, header in enumerate(headers):
        if header.strip().upper() == target:
            return index
    return -1


def _cell(row: list[str], index: int) -> str:
    if index < 0 or index >= len(row):
        return ""
    return row[index].strip()


def _is_day_banner(row: list[str]) -> bool:
    text = [value.strip() for value in row if value.strip()]
    return 0 < len(text) <= 2 and all(DAY_BANNER.match(value) for value in text)


def _weekday(raw: str) -> tuple[str, str]:
    match = SHEET_DATE.search(raw)
    if not match:
        return "", ""
    month, day, year = (int(match.group(1)), int(match.group(2)), int(match.group(3)))
    try:
        stamp = datetime(year, month, day)
    except ValueError:
        return "", ""
    return f"{DAY_NAMES[stamp.weekday()]} · {month}/{day}", stamp.date().isoformat()


def vehicles_from_rows(rows: list[tuple[list[str], bool]]) -> list[dict]:
    header_at = next(
        (index for index, (cells, _) in enumerate(rows) if cells and cells[0].strip().upper() == "VIN"),
        -1,
    )
    if header_at < 0:
        return []
    headers = rows[header_at][0]
    vin_col = _header_index(headers, "VIN")
    year_col = _header_index(headers, "YR")
    make_col = _header_index(headers, "MAKE")
    model_col = _header_index(headers, "MODEL")
    trim_col = _header_index(headers, "TRIM")
    color_col = _header_index(headers, "COLOR")
    odo_col = _header_index(headers, "ODO")
    lane_col = _header_index(headers, "LN")
    lot_col = _header_index(headers, "LOT")
    auction_col = _header_index(headers, "AUCTION")
    cars = []
    banner = ""
    for index in range(header_at + 1, len(rows)):
        cells, dead = rows[index]
        if _is_day_banner(cells):
            banner = " ".join(value.strip() for value in cells if value.strip())
            continue
        vin = _cell(cells, vin_col).upper()
        year_raw = _cell(cells, year_col)
        if len(vin) < 6 or not re.fullmatch(r"\d{4}", year_raw):
            continue
        date_raw = cells[9].strip() if len(cells) > 9 else ""
        label, sort_key = _weekday(date_raw)
        location = _cell(cells, auction_col)
        channel = _cell(cells, auction_col - 1) if auction_col > 0 else ""
        if location and not re.fullmatch(r"simulcast", location, re.I):
            auction = location
        elif channel and not re.fullmatch(r"simulcast", channel, re.I):
            auction = channel
        else:
            auction = location or channel
        odo_raw = _cell(cells, odo_col).replace(",", "")
        try:
            odometer = f"{int(float(odo_raw)):,}" if odo_raw else ""
        except ValueError:
            odometer = _cell(cells, odo_col)
        struck = any("\u0336" in value for value in (vin, _cell(cells, make_col), _cell(cells, model_col)))
        cars.append(
            {
                "vin": vin,
                "year": int(year_raw),
                "make": _cell(cells, make_col),
                "model": _cell(cells, model_col),
                "trim": _cell(cells, trim_col),
                "color": _cell(cells, color_col),
                "odometer": odometer,
                "lane": _cell(cells, lane_col),
                "lot": _cell(cells, lot_col),
                "auction": auction,
                "day_label": label or banner or "Undated",
                "sort_key": sort_key or "9999-99-99",
                "dead": dead or struck,
            }
        )
    return cars


def _xlsx_rows(data: bytes) -> list[tuple[list[str], bool]]:
    book = load_workbook(BytesIO(data), data_only=True)
    sheet = book.active
    rows = []
    for row in sheet.iter_rows():
        cells = []
        dead = False
        for cell in row:
            value = "" if cell.value is None else str(cell.value).strip()
            cells.append(value)
            if value and cell.font is not None and cell.font.strike:
                dead = True
        if any(cells):
            rows.append((cells, dead))
    return rows


def _xls_rows(data: bytes) -> list[tuple[list[str], bool]]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    except xlrd.XLRDError:
        book = xlrd.open_workbook(file_contents=data)
    sheet = book.sheet_by_index(0)
    rows = []
    for row_index in range(sheet.nrows):
        cells = []
        dead = False
        for col_index in range(sheet.ncols):
            cell = sheet.cell(row_index, col_index)
            value = "" if cell.value in (None, "") else str(cell.value).strip()
            if value.endswith(".0") and value[:-2].isdigit():
                value = value[:-2]
            cells.append(value)
            font = None
            if getattr(book, "formatting_info", False) and hasattr(book, "font_list"):
                try:
                    xf = book.xf_list[cell.xf_index]
                    font = book.font_list[xf.font_index]
                except (IndexError, AttributeError):
                    font = None
            if value and font is not None and getattr(font, "struck_out", 0):
                dead = True
        if any(cells):
            rows.append((cells, dead))
    return rows


def _csv_rows(data: bytes) -> list[tuple[list[str], bool]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(StringIO(text))
    return [([cell.strip() for cell in row], False) for row in reader if any(cell.strip() for cell in row)]


def parse_boxes_file(data: bytes, filename: str) -> list[dict]:
    lower = filename.lower()
    if data[:2] == b"PK" or lower.endswith(".xlsx"):
        rows = _xlsx_rows(data)
    elif data[:4] == b"\xd0\xcf\x11\xe0" or lower.endswith(".xls"):
        rows = _xls_rows(data)
    else:
        rows = _csv_rows(data)
    return vehicles_from_rows(rows)
